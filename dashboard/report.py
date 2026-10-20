"""KPKM report (Word, editable): the advisory for one granary, or an overview of all granaries, ready for decision, planning and a final conclusion.

    build_single(D, case, acts, ai_text)  -> bytes of a .docx for one granary
    build_all(D, cases)                    -> bytes of a .docx covering every granary, ranked for attention

Sections of the single-granary report
    1 Summary            key figures and a plain-language summary
    2 Interventions      what to do, who, by when, why (priority labelled Urgent / Soon / Plan ahead / Routine)
    3 AI summary         only when the server has an AI key (checked against the evidence)
    4 Decision and planning   a table KPKM fills in: decision, responsible officer, target date, status, remarks
    5 Final conclusion   a box and sign-off lines for KPKM
    6 Basis and limits   confidence, how the numbers were made, sources, and the evidence used
It is a draft decision-support document, not an official recommendation, and says so on page 1.
"""
import io
import re
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import bridge, tracker
from .knowledge import REFERENCES

GREEN, GREY = "1F5C3A", "595959"
FILL = {"Urgent": "F8D7DA", "Soon": "FFE8B3", "Plan ahead": "E2F0D9", "Routine": "EDEDED"}
PAGE_W = 17.0                                             # cm of usable width on A4 with 2 cm margins


# ---------------------------------------------------------------------------------------------- low-level helpers
def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")                          # CLEAR, never SOLID (SOLID renders black)
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _borders(table, color="BFBFBF", size=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(size))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _fixed(table, widths):
    """fixed layout with explicit column widths (cm), so Word and other viewers agree"""
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblPr = table._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = Cm(w)
    for col, w in zip(table.columns, widths):
        col.width = Cm(w)


def _row_flags(row, header=False, min_height=None):
    trPr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)
    if header:
        h = OxmlElement("w:tblHeader")
        trPr.append(h)
    if min_height:
        row.height = Cm(min_height)
        trPr_h = trPr.find(qn("w:trHeight"))
        if trPr_h is not None:
            trPr_h.set(qn("w:hRule"), "atLeast")


def _cell(cell, text, bold=False, size=9, color=None, align=None, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    r = p.add_run(str(text))
    r.bold, r.italic, r.font.size = bold, italic, Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p


def _table(doc, header, rows, widths, fills=None, size=9, min_height=None, header_fill=GREEN):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    _borders(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        _cell(c, h, bold=True, size=size, color="FFFFFF")
        _shade(c, header_fill)
    _row_flags(t.rows[0], header=True)
    for r, row in enumerate(rows, 1):
        for i, v in enumerate(row):
            _cell(t.rows[r].cells[i], v, size=size)
            if fills and fills.get((r - 1, i)):
                _shade(t.rows[r].cells[i], fills[(r - 1, i)])
        _row_flags(t.rows[r], min_height=min_height)
    _fixed(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def _field(run, instr):
    for kind, text in (("begin", None), (None, instr), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
        else:
            f = OxmlElement("w:instrText")
            f.set(qn("xml:space"), "preserve")
            f.text = text
        run._r.append(f)


def _para(doc, text="", bold=False, size=10.5, color=None, italic=False, after=6, align=None, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.bold, r.italic, r.font.size = bold, italic, Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def _bullet(doc, text, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    _inline(p, text, size)
    return p


def _inline(p, text, size=10):
    """**bold** segments inside a line"""
    for i, seg in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if seg:
            r = p.add_run(seg)
            r.bold = i % 2 == 1
            r.font.size = Pt(size)


def _markdown(doc, text):
    """small markdown subset (headings, bullets, **bold**) for the AI summary"""
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("#"):
            _para(doc, s.lstrip("# ").strip(), bold=True, size=10.5, color=GREEN, after=2)
        elif s[:2] in ("- ", "* "):
            _bullet(doc, s[2:])
        else:
            p = _para(doc, after=4)
            _inline(p, s)


def _h1(doc, text, new_page=False):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.page_break_before = new_page
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    return h


def _h2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    return h


def _new_doc(title_for_header):
    doc = Document()
    s = doc.sections[0]
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)           # A4
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.0)
    st = doc.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    st["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for name, size in (("Heading 1", 15), ("Heading 2", 12), ("Title", 22)):
        f = st[name].font
        f.name, f.size, f.bold = "Calibri", Pt(size), True
        f.color.rgb = RGBColor.from_string(GREEN)
        st[name]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    for nm in ("Header", "Footer"):                              # field results (page numbers) take their size from the style
        st[nm].font.size = Pt(8.5)
        st[nm].font.color.rgb = RGBColor.from_string(GREY)
    hp = s.header.paragraphs[0]
    r = hp.add_run(f"KPKM  |  Paddy granary advisory  |  {title_for_header}")
    r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(GREY)
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"Draft decision support, generated {date.today():%d %B %Y}   |   Page ")
    r.font.size, r.font.color.rgb = Pt(8.5), RGBColor.from_string(GREY)
    r2 = fp.add_run()
    r2.font.size = Pt(8.5)
    _field(r2, "PAGE")
    r3 = fp.add_run(" of ")
    r3.font.size, r3.font.color.rgb = Pt(8.5), RGBColor.from_string(GREY)
    r4 = fp.add_run()
    r4.font.size = Pt(8.5)
    _field(r4, "NUMPAGES")
    return doc


def _notice(doc):
    t = doc.add_table(rows=1, cols=1)
    _borders(t, color="E0B84A")
    c = t.rows[0].cells[0]
    _shade(c, "FFF4D6")
    _cell(c, "DRAFT decision support. It is built from satellite and weather data and a statistical forecast. Satellite greenness cannot "
             "diagnose a cause, so each intervention is something to check, ask or prepare, not a confirmed problem. Review with the "
             "granary authority and an agronomist before acting.", size=9)
    _fixed(t, [PAGE_W])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _title(doc, title, subtitle, meta):
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p.paragraph_format.space_after = Pt(2)
    _para(doc, subtitle, bold=True, size=13, color=GREY, after=2)
    _para(doc, meta, size=9.5, color=GREY, after=8)


# ---------------------------------------------------------------------------------------------- shared blocks
def _key_figures(doc, D, case):
    rows = []
    if case["has_sat"]:
        rows += [("Crop status", f"{case['status'].upper()}  ({case['below']} of {case['checked']} recent periods below range)"),
                 ("Stage in 30 days", f"{case['stage']} (estimated)"),
                 ("NDVI now and in 30 days", f"{case['last_ndvi']:.2f} to {case['f30']:.2f}  (range {case['lo30']:.2f} to {case['hi30']:.2f})"),
                 ("NDVI trend", case["trend"]),
                 ("Data as of", f"{case['as_of']}  ({case['stale_days']} days before this report)")]
    rows.append(("Confidence", case["confidence"].upper()))
    if "res_class" in case:
        rows.append(("Resilience", f"{case['res_class']} third  (index {case['res_index']:.0f} of 100, relative to the other granaries)"))
    if "yield_last" in case:
        rows.append((f"Reported yield {case['yield_year']}", f"{case['yield_last']:,.0f} kg/ha"))
    if "outlook" in case:
        rows.append((f"Outlook {case['outlook_year']}", f"{case['outlook']:,.0f} kg/ha  (range {case['outlook_low']:,.0f} to {case['outlook_high']:,.0f})"))
    t = doc.add_table(rows=len(rows), cols=2)
    _borders(t)
    for i, (k, v) in enumerate(rows):
        _cell(t.rows[i].cells[0], k, bold=True, size=9.5)
        _shade(t.rows[i].cells[0], "EAF2EC")
        _cell(t.rows[i].cells[1], v, size=9.5)
        _row_flags(t.rows[i])
    _fixed(t, [4.6, PAGE_W - 4.6])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _interventions(doc, acts):
    if not acts:
        _para(doc, "No rule matches this granary's current situation.", italic=True)
        return
    rows, fills = [], {}
    for i, a in enumerate(acts):
        lab = bridge.priority_label(a[7])
        rows.append((i + 1, lab, a[0], a[1], a[2], a[3]))
        fills[(i, 1)] = FILL[lab]
    _table(doc, ["No", "Priority", "Intervention", "Responsible", "By when", "Why"], rows,
           [0.9, 1.7, 6.0, 3.0, 2.3, 3.1], fills=fills, size=8.5)


def _fmt_date(s):
    try:
        return date.fromisoformat(str(s)[:10]).strftime("%d %b %Y") if s else ""
    except ValueError:
        return str(s)


def _recorded_cells(rec):
    """decision, responsible, date, status, remarks for one saved decision (blank if nothing was decided)"""
    if rec is None or rec.decision in ("Pending", ""):
        return "", "", "", "", ""
    parts = []
    if rec.kpkm_action:
        parts.append(("KPKM will: " if rec.decision == "Modify" else "Next step: ") + rec.kpkm_action)
    if rec.reason:
        parts.append(("Reason: " if rec.decision in ("Reject", "Defer") else "Remarks: ") + rec.reason)
    return rec.decision, rec.responsible, _fmt_date(rec.target_date), rec.status, "  ".join(parts)


def _decision_table(doc, acts, recorded=None, per_granary=None):
    """table of KPKM decisions: prefilled from what was recorded in the app, otherwise blank to fill in"""
    recorded = recorded or {}
    rows = []
    for a in acts:
        short = a[0] if len(a[0]) <= 150 else a[0][:147].rstrip() + "..."
        head = f"[{bridge.priority_label(a[7])}] {short}" if per_granary is None else f"{per_granary}: {short}"
        d, who, when, st_, rem = _recorded_cells(recorded.get(a[4]))
        rows.append((head, d, who, when, st_, rem))
    if not rows:
        rows = [("", "", "", "", "", "")]
    rows += [("", "", "", "", "", "")] * 2                 # spare rows for additions
    _table(doc, ["Intervention", "Decision (Accept / Modify / Reject / Defer)", "Responsible officer", "Target date", "Status", "Remarks"],
           rows, [5.2, 3.0, 2.6, 1.9, 1.6, 2.7], size=8.5, min_height=1.2)


def _conclusion(doc, text=""):
    _para(doc, "Overall decision, planning notes and final conclusion by KPKM:", size=10, after=3)
    t = doc.add_table(rows=1, cols=1)
    _borders(t, color="7F7F7F", size=6)
    _cell(t.rows[0].cells[0], text, size=10)
    _row_flags(t.rows[0], min_height=4.0)
    _fixed(t, [PAGE_W])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    t2 = doc.add_table(rows=2, cols=4)
    _borders(t2)
    for i, h in enumerate(("Name", "Designation", "Date", "Signature")):
        _cell(t2.rows[0].cells[i], h, bold=True, size=9, color="FFFFFF")
        _shade(t2.rows[0].cells[i], GREEN)
        _cell(t2.rows[1].cells[i], "")
    _row_flags(t2.rows[0])
    _row_flags(t2.rows[1], min_height=1.4)
    _fixed(t2, [4.8, 4.4, 3.0, 4.8])


def _sources(doc, acts_list):
    used = sorted({s.strip() for acts in acts_list for a in acts for s in a[6].split(",") if s.strip() in REFERENCES})
    if not used:
        return
    _h2(doc, "Sources behind the interventions")
    for k in used:
        kind, cite, url = REFERENCES[k]
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{k} ({kind}): ")
        r.bold, r.font.size = True, Pt(9)
        r = p.add_run(f"{cite} {url}")
        r.font.size = Pt(9)
    _para(doc, "Several sources are summaries (news reports, a library record, a practitioner site) and not the original manuals. Check "
               "the originals and the publishers' terms before quoting them in an official document.", italic=True, size=8.5, color=GREY)


def _method_notes(doc, D):
    dyn = getattr(D, "yield_dyn", None)
    err = ""
    if dyn is not None and "ar1_transfer" in dyn.columns:
        v = dyn.set_index("horizon").ar1_transfer
        err = " Back-tested error of the yield outlook: " + ", ".join(f"about {v[h]:.0f}% at {h} year{'s' if h > 1 else ''}" for h in v.index) + "."
    for t in (
        "**Crop status** compares the satellite greenness (NDVI) with the range a trained forecast expects for this stage. 'Behind' means "
        "it was below that range in at least 2 of the last 4 ten-day periods. It says the crop looks different from expected, not why.",
        "**Resilience** is measured through the same forecast: how often and how deeply greenness falls below the expected range, how fast "
        "it recovers, and how steady the yearly peak is. It is relative to these granaries and has not been shown to predict yield.",
        "**Yield outlook** uses transfer learning of yield behaviour from the other granaries. It beats training on one granary alone but is "
        "no better than 'same as last year'." + err,
        "**Stage** is estimated from the planting calendar, not observed in the field. **Interventions** come from draft rules that an "
        "agronomist has not yet reviewed."):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        _inline(p, t, 9.5)


def _score(c):
    return (c["below"] * 10 + c["below_share"]) if c["has_sat"] else -1


# ---------------------------------------------------------------------------------------------- single granary
def _has_decisions(recorded):
    return any(getattr(r, "decision", "") not in ("", "Pending", "Noted") for r in (recorded or {}).values())


def build_single(D, case, acts, ai_text=None, recorded=None):
    g = case["granary"]
    recorded = recorded or {}
    doc = _new_doc(g)
    _title(doc, "Granary Advisory and Intervention Report", g,
           (case["admin"] + "   |   " if case["admin"] else "") + f"Prepared for KPKM   |   {date.today():%d %B %Y}")
    _notice(doc)
    _h1(doc, "1. Summary")
    _key_figures(doc, D, case)
    for s in bridge.summary_text(D, case):
        _para(doc, s, after=4)
    _h1(doc, "2. Recommended interventions")
    _para(doc, "Ordered from most to least urgent. 'Responsible' is the suggested lead; KPKM can reassign it in section 4.", size=9.5,
          color=GREY, after=4)
    _interventions(doc, acts)
    n = 3
    if ai_text:
        _h1(doc, f"{n}. AI summary")
        _para(doc, "Written by an AI model from the evidence and the interventions above only. Check it before use.", italic=True, size=9, color=GREY)
        _markdown(doc, ai_text)
        n += 1
    _h1(doc, f"{n}. Decision and planning " + ("(recorded in the app; KPKM to confirm)" if _has_decisions(recorded)
                                               else "(to be completed by KPKM)"), new_page=True)
    _decision_table(doc, acts, recorded)
    _h1(doc, f"{n + 1}. Final conclusion")
    ov = recorded.get(tracker.OVERALL)
    _conclusion(doc, ov.kpkm_action if ov is not None else "")
    _h1(doc, f"{n + 2}. Basis, confidence and limits")
    p = _para(doc, after=4)
    _inline(p, f"**Confidence: {case['confidence']}.** " + "; ".join(case["conf_reasons"]).capitalize() + ".", 10)
    _method_notes(doc, D)
    _sources(doc, [acts])
    _h2(doc, "Evidence used")
    for i, t in bridge.evidence(D, case):
        _bullet(doc, f"[{i}] {t}", 9)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------------- all granaries
def build_all(D, cases, cur=None):
    if cur is None and hasattr(D, "DATA_DIR"):
        cur = tracker.current(tracker.load(D.DATA_DIR))           # the latest saved decision of every item
    ranked = sorted(cases.values(), key=_score, reverse=True)
    acts = {c["granary"]: bridge.interventions(c, D.irules) for c in ranked}
    doc = _new_doc("all granaries")
    _title(doc, "Paddy Granaries: Priority Overview and Advisory", f"{len(ranked)} granaries",
           f"Prepared for KPKM   |   {date.today():%d %B %Y}   |   ranked by recent shortfall against the expected greenness range")
    _notice(doc)
    _h1(doc, "1. Priority overview")
    rows = []
    for i, c in enumerate(ranked, 1):
        a = acts[c["granary"]]
        first = a[0] if a else None
        rows.append((i, c["granary"], c.get("status", "no satellite"), c["confidence"], c.get("res_class", "-"),
                     (c.get("stage", "-") or "-").split(" / ")[0],
                     (f"{bridge.priority_label(first[7])}: " + (first[0][:70] + ("..." if len(first[0]) > 70 else ""))) if first else "-"))
    _table(doc, ["#", "Granary", "Crop status", "Confidence", "Resilience", "Stage in 30 days", "First intervention"], rows,
           [0.8, 3.1, 1.9, 1.9, 1.7, 2.1, 5.5], size=8.5)
    for c in ranked:
        _h1(doc, f"{c['granary']}")
        if c["admin"]:
            _para(doc, c["admin"], size=9, color=GREY, after=3)
        for s in bridge.summary_text(D, c):
            _para(doc, s, size=10, after=3)
        _h2(doc, "Interventions")
        _interventions(doc, acts[c["granary"]])
    recs = {c["granary"]: (tracker.recorded_for(cur, c["granary"], c.get("as_of", "")) if cur is not None else {}) for c in ranked}
    _h1(doc, "Decision and planning " + ("(recorded in the app; KPKM to confirm)" if any(_has_decisions(r) for r in recs.values())
                                         else "(to be completed by KPKM)"), new_page=True)
    rows = []
    for c in ranked:
        for a in acts[c["granary"]][:3]:
            short = a[0] if len(a[0]) <= 80 else a[0][:77].rstrip() + "..."
            d, who, when, st_, rem = _recorded_cells(recs[c["granary"]].get(a[4]))
            rows.append((f"{c['granary']}: [{bridge.priority_label(a[7])}] {short}", d, who, when, st_, rem))
    rows += [("", "", "", "", "", "")] * 3
    _table(doc, ["Granary and intervention", "Decision (Accept / Modify / Reject / Defer)", "Responsible officer", "Target date", "Status", "Remarks"],
           rows, [5.2, 3.0, 2.6, 1.9, 1.6, 2.7], size=8.5, min_height=1.2)
    _para(doc, "The table lists the top three interventions per granary; the full list is in each granary's section above.", italic=True,
          size=8.5, color=GREY)
    _h1(doc, "Final conclusion")
    ovs = [f"{g}: {r[tracker.OVERALL].kpkm_action}" for g, r in recs.items() if tracker.OVERALL in r and r[tracker.OVERALL].kpkm_action]
    _conclusion(doc, "\n".join(ovs))
    _h1(doc, "Basis, confidence and limits")
    _method_notes(doc, D)
    _sources(doc, list(acts.values()))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
